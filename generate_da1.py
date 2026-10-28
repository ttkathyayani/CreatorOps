import os
import docx
from docx.shared import Inches, Pt, RGBColor
from build_reports_docx import (
    init_styled_doc, add_title_block, add_heading_1, add_heading_2, add_heading_3,
    add_body_p, add_callout, add_styled_table, add_figure_with_caption
)

def build_da1():
    print("Generating DA1 (Review 1 Report)...")
    doc = init_styled_doc("CreatorOps SaaS", "BCSE408L", "DA1 / Review 1 Deliverables")
    
    # Title Block
    add_title_block(
        doc,
        main_title="CreatorOps — Cloud-Native Content Operations Platform",
        subtitle="Digital Assignment 1 / Project Review 1 Report: Problem Statement, System Scope, Non-Vague Tech Stack Specifications, Requirements Analysis, UML Use Case Specifications, UI Wireframes & Initial SRS",
        stage_label="Digital Assignment 1 (Review 1 Deliverable)",
        candidate_name="Kathyayani T T"
    )
    
    # 1. Executive Summary & Problem Context
    add_heading_1(doc, "1. Executive Summary & Problem Context")
    add_body_p(doc, 
               "CreatorOps is an enterprise-grade, cloud-native Software-as-a-Service (SaaS) application engineered for the modern Creator Economy and Marketing Agency domain. "
               "As social media campaigns expand across platforms (Instagram, YouTube, TikTok, LinkedIn), creative teams and agencies face crippling operational chaos due to scattered feedback, "
               "lack of version control, premature client publishing, and zero structured audit trails.")
    
    add_callout(doc, 
                "Core Value Proposition: CreatorOps replaces fragmented messaging channels with 'One campaign, One workspace, One approval trail.' "
                "It enforces a strict 8-state workflow machine, dual-tier approval gating (Agency Manager & Brand Client), automated script quality audits, "
                "centralized cloud media asset storage, and autonomous Public Social Media Competitor Benchmarking (YouTube, TikTok, Instagram).",
                title="SaaS Operational Mandate", callout_type="info")
    
    # 2. Problem Statement & Existing System Deficiencies
    add_heading_1(doc, "2. Problem Statement & Existing System Deficiencies")
    add_body_p(doc, 
               "In existing marketing agency workflows, content creation, collaboration, and client sign-offs rely on disconnected consumer tools such as WhatsApp groups, Slack messages, Google Drive folders, and Notion boards. "
               "Furthermore, marketing teams lack visibility into public competitor performance, producing content blindly without data on what hooks or formats drive engagement on social media platforms. "
               "This ad-hoc methodology leads to systemic operational failures:")
    
    prob_headers = ["Friction Area", "Existing Ad-Hoc Method", "Failure Mode / Business Risk", "CreatorOps Cloud Solution"]
    prob_rows = [
        ["Feedback & Revisions", "Scattered WhatsApp, Slack, & Email threads", "Conflicting feedback, missed comments, 43% rework delays", "Contextual threaded comments locked to exact content versions"],
        ["Version Integrity", "Files named 'video_final_v2_edit.mp4' on Drive", "Accidental publishing of outdated revisions & overwritten edits", "Automated immutable version snapshotting with visual diff viewer"],
        ["Approval Governance", "Verbal or informal chat 'looks good' messages", "Premature posting without client sign-off; breach of contract", "Two-tier cryptographic approval gates (Internal Manager -> Client)"],
        ["Media Asset Control", "Scattered personal cloud links with expired URLs", "Missing creative assets, broken download links, asset loss", "Centralized Amazon S3 cloud object storage with presigned URLs"],
        ["Brand Compliance", "Manual proofreading of scripts and hooks", "Publishing non-compliant claims, banned terms, poor retention", "Automated AI script quality, pacing, hook placement, & safety audits"],
        ["Market Competitor Gap", "Zero tracking of competitor post reach or viral hooks", "Blind content production; lower reach; wasted ad spend", "Automated Public Social Ingestion & Head-to-Head Benchmarking Matrix"]
    ]
    add_styled_table(doc, prob_headers, prob_rows, [Inches(1.2), Inches(1.8), Inches(1.8), Inches(1.7)])
    
    # 3. Proposed Cloud-Native System
    add_heading_1(doc, "3. Proposed System Architecture & Core Innovations")
    add_body_p(doc, 
               "The proposed CreatorOps platform introduces a multi-tenant, cloud-native architecture designed to eliminate operational friction through five foundational innovations:")
    
    add_body_p(doc, "Multi-Tenant Agency Workspaces: Provides complete tenant data isolation, enabling agencies to manage multiple separate brand clients and creator teams within segregated security boundaries.", bold_prefix="•")
    add_body_p(doc, "8-State Sequential Approval State Machine: Enforces a programmatic state transition model where content deliverables cannot bypass internal agency checks or client sign-offs.", bold_prefix="•")
    add_body_p(doc, "Automated Script & Content Quality Engine: Real-time analysis of script text for hook effectiveness, estimated speech duration, readability score, and brand safety compliance.", bold_prefix="•")
    add_body_p(doc, "Cloud Object Storage Abstraction: Pluggable storage architecture supporting local development and production cloud object storage (Amazon S3 / MinIO) for high-resolution video and image assets.", bold_prefix="•")
    add_body_p(doc, "Public Social Ingestion & Competitor Benchmarking: Autonomous telemetry collection aggregating public video views, likes, comments, and engagement rates across YouTube, TikTok, and Instagram to calculate market-relative performance and AI content optimization tips.", bold_prefix="•")
    
    # 4. Non-Vague Technology Stack Specifications
    add_heading_1(doc, "4. Non-Vague Technology Stack Specifications")
    add_body_p(doc, 
               "To satisfy enterprise cloud standards and academic rigor, CreatorOps employs an explicit, modern technology stack with pinned component versions, concrete runtime specifications, and dedicated cloud tiers:")
    
    tech_headers = ["Layer / Subsystem", "Technology & Exact Version", "Runtime / Protocol / Port", "Architectural Role & Scope"]
    tech_rows = [
        ["Frontend UI & SPA", "React 19.0.0, Vite 5.3.1, TypeScript 5.4.5", "Browser DOM / Node.js 20.14 LTS", "Responsive dark-mode UI, reactive state, client routing, zero-bundle overhead"],
        ["CSS & Design System", "Tailwind CSS v4.0.0, Lucide React 0.395", "PostCSS 8.5 / CSS Variables", "Dark slate styling (#0f172a, #1e293b, #4f46e5), high-contrast accessibility"],
        ["Backend REST Engine", "FastAPI 0.111.0, Starlette 0.37.2", "Python 3.12.2 / ASGI (Uvicorn 0.30)", "Asynchronous route handling, OpenAPI 3.1 generation, dependency injection"],
        ["Data Serialization", "Pydantic v2.7.4, pydantic-settings 2.2.1", "Native Rust-backed core validation", "Strict JSON schema enforcement, request validation, environment configuration"],
        ["Database ORM", "SQLAlchemy 2.0.30, Alembic 1.13.1", "Python / Connection Pooling (Pool 10)", "Normalized 3NF relational mapping, declarative models, automated migrations"],
        ["Relational Database", "PostgreSQL 14.23 (Local port 5433 / Neon Cloud)", "TCP Port 5432/5433, SSL required", "ACID transactional guarantees, relational foreign keys, multi-tenant schemas"],
        ["Cloud Object Storage", "Amazon S3 API / MinIO (Boto3 1.34.120)", "HTTPS REST / Presigned URLs", "High-throughput asset storage for 4K video drafts, scripts, and brand guidelines"],
        ["Public Social Scraper", "YouTube Data API v3, HTTPX 0.27, BS4 4.12", "Async HTTPS REST / JSON Ingestion", "Autonomous harvesting of public video views, likes, and competitor metrics"],
        ["Security & Auth", "Passlib 1.7.4 (PBKDF2 SHA-256 390k), Jose 3.3", "HMAC-SHA256 Signed JWT Tokens", "Cryptographic password hashing, RBAC token verification, audit logging"],
        ["Edge Hosting & CI/CD", "Vercel Anycast Edge Platform, GitHub Actions", "Global CDN / Git Hook Deployments", "Zero-downtime atomic rollbacks, sub-200ms global TTFB, CI lint & test pipeline"]
    ]
    add_styled_table(doc, tech_headers, tech_rows, [Inches(1.2), Inches(2.1), Inches(1.6), Inches(1.6)])
    
    # 5. Target Stakeholders & RBAC
    add_heading_1(doc, "5. Target Stakeholders & Role-Based Access Control (RBAC)")
    add_body_p(doc, 
               "To satisfy enterprise security guidelines, CreatorOps establishes strict Role-Based Access Control across five primary stakeholder archetypes:")
    
    rbac_headers = ["Stakeholder Role", "Operational Scope", "Primary Responsibilities", "System Permissions & Access"]
    rbac_rows = [
        ["Content Creator", "Creative Studio & Workspace", "Drafts scripts, uploads raw media, views revision requests, inspects hook recommendations.", "Create & edit assigned content; upload media; view contextual comments."],
        ["Agency Manager", "Campaign & Review Governance", "Allocates campaign budgets, schedules timelines, reviews creator drafts, benchmarks rivals.", "CRUD on campaigns; review content; approve internal drafts; competitor analytics."],
        ["Brand Client", "External Review Portal", "Reviews manager-approved deliverables, provides brand feedback, grants final sign-off.", "View approved deliverables; post timestamped feedback; approve/reject content."],
        ["Market Analyst", "Competitor Intelligence Hub", "Monitors public competitor views, analyzes viral hooks, generates content gap reports.", "Inspect competitor telemetry; filter platforms; trigger ingestion refresh."],
        ["System Administrator", "Platform & Infrastructure", "Manages agency workspaces, provisions cloud databases, monitors API uptime.", "Global tenant configuration; audit log inspection; system health monitoring."]
    ]
    add_styled_table(doc, rbac_headers, rbac_rows, [Inches(1.2), Inches(1.5), Inches(2.1), Inches(1.7)])
    
    # 6. Requirements Specifications (FR & NFR)
    add_heading_1(doc, "6. Software Requirements Specifications (SRS)")
    add_heading_2(doc, "6.1 Functional Requirements Matrix")
    
    fr_headers = ["Req ID", "Requirement Name", "System Trigger & Input", "Expected System Output", "Priority"]
    fr_rows = [
        ["FR-01", "User Authentication & RBAC", "User registration / Login credentials", "Cryptographically hashed PBKDF2 verification; signed JWT token", "Mandatory (High)"],
        ["FR-02", "Multi-Tenant Workspaces", "Create agency workspace & slug", "Isolated workspace database bounds; tenant switcher", "Mandatory (High)"],
        ["FR-03", "Client Brand CRM", "Client profile creation & guidelines", "Persisted brand roster, contact emails, and visual standards", "Mandatory (High)"],
        ["FR-04", "Campaign Operations", "Create campaign, budget, & dates", "Campaign timeline, channel allocation, and objective tracking", "Mandatory (High)"],
        ["FR-05", "Content Studio & Editing", "Title, social channel, & script body", "Saved deliverable with word counter, hook score, & status badge", "Mandatory (High)"],
        ["FR-06", "Immutable Version History", "Creator updates script or assets", "Automatic creation of sequential revision snapshot (Rev 1, Rev 2...)", "Mandatory (High)"],
        ["FR-07", "Cloud Media Storage", "Upload video/image assets (up to 100MB)", "S3 cloud object storage persistence; signed CDN delivery URLs", "Mandatory (High)"],
        ["FR-08", "Two-Tier Approval Gate", "Manager / Client approve/reject action", "State machine transition; audit log record; email notification", "Mandatory (High)"],
        ["FR-09", "Automated Script Auditing", "Trigger AI script review check", "Readability grade, hook rating, speaking duration, safety flag", "Desirable (Med)"],
        ["FR-10", "Operational Analytics", "Aggregated campaign metrics request", "Dashboard reporting on approval SLA, active campaigns, deliverables", "Mandatory (High)"],
        ["FR-11", "Public Social Ingestion & Benchmarking", "Competitor handle or URL (YT/TikTok/IG)", "Public views, likes, engagement % vs benchmark, and AI hook tips", "Mandatory (High)"]
    ]
    add_styled_table(doc, fr_headers, fr_rows, [Inches(0.7), Inches(1.6), Inches(1.8), Inches(1.8), Inches(0.8)])
    
    add_heading_2(doc, "6.2 Non-Functional Requirements (NFR) & Quality Benchmarks")
    nfr_headers = ["Quality Attribute", "Specification & Architectural Target", "Verification Mechanism", "Target Metric"]
    nfr_rows = [
        ["Security", "PBKDF2 SHA-256 (390k rounds) hashing, TLS 1.3 encryption, SQL parameterized queries", "OWASP compliance & automated security audit", "Zero SQLi / XSS vulnerabilities"],
        ["Performance", "Sub-second API response time; static asset edge CDN caching", "FastAPI middleware & CDN telemetry", "< 250ms API p95 latency"],
        ["Availability", "Serverless multi-region Anycast edge deployment; automated zero-downtime rollouts", "Vercel edge status & continuous ping health check", ">= 99.9% Uptime SLA"],
        ["Scalability", "Stateless serverless compute scaling from 0 to N concurrent execution instances", "Locust load testing under peak concurrent load", "Horizontal elastic scaling"],
        ["Reliability", "PostgreSQL ACID transaction isolation, relational foreign key cascade rules", "Alembic schema validations & Pytest test suite", "100% data integrity"],
        ["Usability", "Responsive dark-mode UI optimized for desktop, tablet, and mobile form factors", "Chrome DevTools audits & cross-device validation", ">= 90 Lighthouse score"]
    ]
    add_styled_table(doc, nfr_headers, nfr_rows, [Inches(1.2), Inches(2.2), Inches(1.8), Inches(1.3)])
    
    # 7. Project Scope & System Boundaries
    add_heading_1(doc, "7. Project Scope & System Boundaries")
    add_body_p(doc, 
               "To ensure achievable, high-impact milestones across the semester, project boundaries are clearly categorized into in-scope deliverables and planned future integrations:")
    
    scope_headers = ["Functional Scope Dimension", "In-Scope for CreatorOps Cloud SaaS", "Out-of-Scope (Future Expansion)"]
    scope_rows = [
        ["User & Tenant Boundaries", "Agency multi-tenancy, user registration, login, and 5 RBAC roles.", "Enterprise SSO / SAML federation."],
        ["Campaign Management", "Multi-client roster, campaign scheduling, budget tracking, channel tags.", "Direct programmatic financial invoicing."],
        ["Content Creation & Review", "Script drafting, version snapshots, AI hook auditing, threaded feedback.", "In-browser native video multi-track editing."],
        ["Approval State Machine", "Sequential 8-state transition machine with internal and client approvals.", "Automated legal contract electronic notarization."],
        ["Cloud Storage & Delivery", "Cloud S3/MinIO bucket asset storage with presigned CDN links.", "On-premise physical tape backup archival."],
        ["Publishing & Analytics", "Social media scheduling calendar and aggregated operational KPIs.", "Direct API auto-publishing to Instagram/TikTok."],
        ["Competitor Intelligence", "Public social data aggregation (views, likes, engagement rate) & benchmarking.", "Automated real-time bidding for sponsored ads."]
    ]
    add_styled_table(doc, scope_headers, scope_rows, [Inches(1.5), Inches(2.5), Inches(2.5)])
    
    # 8. UML Use Case Diagram
    add_heading_1(doc, "8. UML Use Case Architecture & Actor Interaction Matrix")
    add_body_p(doc, 
               "Figure 1 illustrates the UML Use Case Diagram for CreatorOps. The system boundary encapsulates the cloud SaaS platform, "
               "delineating the operational capabilities accessible to Creators, Agency Managers, Brand Clients, Market Analysts, and System Administrators.")
    
    add_figure_with_caption(
        doc,
        img_path="docs/figures/use_case_diagram.png",
        figure_number="1",
        title="UML Use Case Diagram for CreatorOps Cloud SaaS Platform",
        description="Comprehensive behavioral use case model illustrating actor interactions, permission boundaries, and functional system use cases."
    )
    
    # 9. User Interface Mockups & Design System
    add_heading_1(doc, "9. User Interface Mockups & Design System")
    add_body_p(doc, 
               "CreatorOps adopts an executive dark-mode design system (#0f172a, #1e293b, #4f46e5) built on Tailwind CSS, featuring high-contrast data visualization, "
               "live health indicators, responsive layout structures, and dedicated competitor analytics displays.")
    
    add_figure_with_caption(
        doc,
        img_path="docs/figures/creatorops_competitor_analytics_live.png",
        figure_number="2",
        title="CreatorOps Live Competitor Intelligence & Public Social Benchmarking Dashboard",
        description="Live executed interface displaying 4 comparative KPI cards (1.42M Total Video Views, 5.8% Engagement Rate, 84 Retention Score, 4.2 Posts/Week), filterable Head-to-Head Performance Matrix, and AI Content Optimization Recommendations."
    )
    
    add_figure_with_caption(
        doc,
        img_path="docs/figures/creatorops_live_dashboard.png",
        figure_number="3",
        title="CreatorOps Executed Application Dashboard & System Health Telemetry",
        description="Real-time running dashboard displaying live PostgreSQL 14 connection (1.67ms query latency), storage status, Alembic 14-table migration schema, and content lifecycle state machine."
    )
    
    add_figure_with_caption(
        doc,
        img_path="docs/figures/approval_studio_ui.jpg",
        figure_number="4",
        title="CreatorOps Approval Studio & Script Review Workspace",
        description="Side-by-side editing interface featuring video asset proofing, script drafting, threaded stakeholder feedback, and one-click approval buttons."
    )
    
    # 10. 8-State Approval Workflow State Machine
    add_heading_1(doc, "10. Core Workflow: 8-State Content Approval State Machine")
    add_body_p(doc, 
               "At the heart of CreatorOps is an immutable 8-state content approval workflow that eliminates premature publishing and guarantees strict sign-off compliance:")
    
    add_figure_with_caption(
        doc,
        img_path="docs/figures/workflow_state_machine.png",
        figure_number="5",
        title="CreatorOps 8-Stage Content Approval Lifecycle State Machine",
        description="Deterministic state progression with dual-tier validation gates, revision loops, and immutable audit logging."
    )
    
    # 11. Initial SRS Summary & Milestone Roadmap
    add_heading_1(doc, "11. Initial SRS Summary & Semester Roadmap")
    add_body_p(doc, 
               "The project execution follows a structured milestone schedule aligned with BCSE408L review criteria:")
    
    road_headers = ["Milestone", "Evaluation Date", "Marks", "Target Deliverables & Verified Outcomes", "Status"]
    road_rows = [
        ["Review 1 (DA1)", "20 August 2026", "10 Marks", "Problem statement, non-vague tech stack, SRS specifications, UML use case diagram, UI wireframes, project scope.", "COMPLETED"],
        ["Review 2 (DA2)", "17 September 2026", "10 Marks", "3-Tier System architecture, cloud edge topology, 3NF database ERD, REST API catalog, initial deployment.", "READY"],
        ["Review 3 (DA3)", "15 October 2026", "10 Marks", "Full SaaS application demonstration, competitor intelligence suite, live public URL, automated test suite, 22-section final report.", "FINAL DEMO"]
    ]
    add_styled_table(doc, road_headers, road_rows, [Inches(1.1), Inches(1.1), Inches(0.8), Inches(2.7), Inches(0.8)])
    
    output_docx = "DA1_CreatorOps_Review1_Report.docx"
    try:
        doc.save(output_docx)
        print(f"Saved {output_docx} successfully.")
    except PermissionError:
        fallback = "DA1_CreatorOps_Review1_Live_Report.docx"
        doc.save(fallback)
        print(f"Note: {output_docx} is locked by Microsoft Word. Saved updated file to {fallback} successfully.")
    
    # Also save in docs/
    try:
        doc.save("docs/DA1_CreatorOps_Review1_Report.docx")
        print("Saved copy to docs/DA1_CreatorOps_Review1_Report.docx")
    except Exception as e:
        pass

if __name__ == "__main__":
    build_da1()
