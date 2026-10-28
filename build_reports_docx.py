import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    kwargs can be top, bottom, left, right with values like:
    {"sz": 12, "val": "single", "color": "CBD5E1"}
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', 'CBD5E1'))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def init_styled_doc(title, header_text, stage_text):
    doc = docx.Document()
    
    # 1 inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.text = f"BCSE408L: Cloud Computing | {title} | {stage_text}"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.style.font.name = 'Segoe UI'
        hp.style.font.size = Pt(8.5)
        hp.style.font.color.rgb = RGBColor(100, 116, 139)
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "Candidate: Kathyayani T T | Project 12: CreatorOps | VIT Chennai"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.style.font.name = 'Segoe UI'
        fp.style.font.size = Pt(8.5)
        fp.style.font.color.rgb = RGBColor(148, 163, 184)
        
    return doc

def add_title_block(doc, main_title, subtitle, stage_label, candidate_name="Kathyayani T T"):
    # Category badge
    p_badge = doc.add_paragraph()
    p_badge.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_badge = p_badge.add_run(f"VELLORE INSTITUTE OF TECHNOLOGY • SCOPE • ACADEMIC YEAR 2026–2027\n{stage_label.upper()}")
    run_badge.font.name = 'Segoe UI'
    run_badge.font.size = Pt(9.5)
    run_badge.font.bold = True
    run_badge.font.color.rgb = RGBColor(79, 70, 229) # Indigo
    
    # Main Title
    p_title = doc.add_paragraph()
    run_title = p_title.add_run(main_title)
    run_title.font.name = 'Segoe UI'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 41, 59) # Slate 800
    p_title.paragraph_format.space_after = Pt(2)
    
    # Subtitle
    p_sub = doc.add_paragraph()
    run_sub = p_sub.add_run(subtitle)
    run_sub.font.name = 'Segoe UI'
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(71, 85, 105) # Slate 600
    p_sub.paragraph_format.space_after = Pt(12)
    
    # Metadata Card Table
    table = doc.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    meta_items = [
        ("Course Code & Title", "BCSE408L — Cloud Computing & Full Stack App Dev"),
        ("Candidate Name", candidate_name),
        ("Project Title", "Project 12: CreatorOps (Marketing SaaS)"),
        ("Live Deployment URL", "https://creatorops-rho.vercel.app"),
        ("GitHub Repository", "https://github.com/ttkathyayani/CreatorOps"),
        ("Evaluation Review", stage_label)
    ]
    
    col_widths = [Inches(2.2), Inches(2.2), Inches(2.1)]
    for i, (k, v) in enumerate(meta_items):
        r = i // 3
        c = i % 3
        cell = table.cell(r, c)
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        set_cell_border(cell, 
                        top={"sz": 4, "val": "single", "color": "E2E8F0"},
                        bottom={"sz": 4, "val": "single", "color": "E2E8F0"},
                        left={"sz": 4, "val": "single", "color": "E2E8F0"},
                        right={"sz": 4, "val": "single", "color": "E2E8F0"})
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        rk = p.add_run(k + "\n")
        rk.font.name = 'Segoe UI'
        rk.font.size = Pt(8)
        rk.font.bold = True
        rk.font.color.rgb = RGBColor(100, 116, 139)
        
        rv = p.add_run(v)
        rv.font.name = 'Segoe UI'
        rv.font.size = Pt(9)
        rv.font.bold = True
        rv.font.color.rgb = RGBColor(15, 23, 42)
        
    for row in table.rows:
        for j, w in enumerate(col_widths):
            row.cells[j].width = w
            
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

def add_heading_1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 58, 138) # Navy / Indigo
    return h

def add_heading_2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(51, 65, 85) # Slate 700
    return h

def add_heading_3(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(2)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = RGBColor(79, 70, 229) # Indigo
    return h

def add_body_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = 'Segoe UI'
        r_pre.font.size = Pt(9.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_callout(doc, text, title=None, callout_type="info"):
    # Types: info (indigo), success (green), warning (amber)
    colors = {
        "info": ("4F46E5", "EEF2FF"),
        "success": ("10B981", "ECFDF5"),
        "warning": ("F59E0B", "FFFBEB"),
    }
    border_color, bg_color = colors.get(callout_type, ("4F46E5", "EEF2FF"))
    
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    set_cell_border(cell,
                    left={"sz": 24, "val": "single", "color": border_color},
                    top={"sz": 0, "val": "none", "color": "FFFFFF"},
                    bottom={"sz": 0, "val": "none", "color": "FFFFFF"},
                    right={"sz": 4, "val": "single", "color": "E2E8F0"})
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if title:
        rt = p.add_run(f"★ {title}\n")
        rt.font.name = 'Segoe UI'
        rt.font.size = Pt(9.5)
        rt.font.bold = True
        rt.font.color.rgb = RGBColor(30, 41, 59)
    
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(51, 65, 85)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(6)

def add_styled_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], "1E3A8A") # Dark Navy
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        set_cell_border(hdr_cells[i],
                        top={"sz": 6, "val": "single", "color": "1E3A8A"},
                        bottom={"sz": 12, "val": "single", "color": "3B82F6"},
                        left={"sz": 4, "val": "single", "color": "2563EB"},
                        right={"sz": 4, "val": "single", "color": "2563EB"})
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.name = 'Segoe UI'
            r.font.size = Pt(8.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            
    # Data Rows
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_col = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=120, right=120)
            set_cell_border(row_cells[c_idx],
                            top={"sz": 4, "val": "single", "color": "E2E8F0"},
                            bottom={"sz": 4, "val": "single", "color": "E2E8F0"},
                            left={"sz": 4, "val": "single", "color": "E2E8F0"},
                            right={"sz": 4, "val": "single", "color": "E2E8F0"})
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Segoe UI'
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(30, 41, 59)
                
    if col_widths:
        for row in table.rows:
            for j, w in enumerate(col_widths):
                row.cells[j].width = w
                
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(8)
    return table

def add_figure_with_caption(doc, img_path, figure_number, title, description, width=Inches(6.2)):
    if not os.path.exists(img_path):
        print(f"Warning: Image {img_path} not found!")
        return
        
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(3)
    p_img.paragraph_format.keep_with_next = True
    
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=width)
    
    # Caption
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(12)
    
    r_num = p_cap.add_run(f"Figure {figure_number}: {title}\n")
    r_num.font.name = 'Segoe UI'
    r_num.font.size = Pt(9)
    r_num.font.bold = True
    r_num.font.color.rgb = RGBColor(30, 41, 59)
    
    r_desc = p_cap.add_run(description)
    r_desc.font.name = 'Segoe UI'
    r_desc.font.size = Pt(8)
    r_desc.font.italic = True
    r_desc.font.color.rgb = RGBColor(100, 116, 139)

print("Base DOCX builder utilities configured successfully.")
